#!/usr/bin/env python3
"""Read exported XLSX caches/markup and prove ledger determinism. No authoring."""
import copy
import contextlib
import hashlib
import importlib.util
import io
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/architecture/12-decisions/decision-packages/current/evaluation/common-qa-v4'
FILES=OUT/'artifacts/outputs/01a15dfe-v4'


def main():
    data=json.loads((OUT/'artifacts/ledger.json').read_text())
    sheetpath=FILES/'VIA-QA-common-v4.xlsx'
    values=openpyxl.load_workbook(sheetpath,data_only=True,read_only=True)
    formulas=openpyxl.load_workbook(sheetpath,data_only=False,read_only=True)
    meta=json.loads((FILES/'render-verification.json').read_text())
    assert values.sheetnames==meta['sheets'] and len(values.sheetnames)==10
    assert meta['ledger_sha256']==hashlib.sha256((OUT/'artifacts/ledger.json').read_bytes()).hexdigest()
    keys=['qa01','qa02','qa03','qa04','qa05','qa06','qa07','first_voice']
    count=0
    for cells in values['전체비교'].iter_rows(min_row=6,max_row=15,min_col=3,max_col=11,values_only=True):
        expected=data['summary'][cells[0]]
        for key,v in zip(keys,cells[1:]):
            if expected[key] is None:assert v=='미산출'
            else:assert math.isclose(v,expected[key],abs_tol=1e-7)
            count+=1
    for sh in values:
        for row in sh:
            for c in row:assert c.data_type!='e',(sh.title,c.coordinate,c.value)
    formula_count=sum(c.data_type=='f' for sh in formulas for row in sh for c in row)
    assert formula_count==meta['formula_count']
    for name,col,last in [('지칭판정',5,236),('의미판정',4,173)]:
        assert all(row[0].value is None for row in values[name].iter_rows(min_row=6,max_row=last,min_col=col,max_col=col))
    values.close();formulas.close()
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(sheetpath) as z:
        assert z.testzip() is None
        # Reject Excel errors and confirm formulas have cached values in export.
        cached=0
        for name in z.namelist():
            if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
                tree=ET.fromstring(z.read(name))
                for c in tree.findall('.//s:c',ns):
                    assert c.get('t')!='e'
                    if c.find('s:f',ns) is not None:
                        if c.find('s:v',ns) is None and c.find('s:is',ns) is None:
                            # Formula result "" may omit an OpenXML cache value.
                            assert (name=='xl/worksheets/sheet7.xml' and c.get('r','').startswith('F')) or (name=='xl/worksheets/sheet8.xml' and c.get('r','').startswith('E')),(name,c.get('r'))
                        cached+=1
        assert cached==formula_count
    checker_spec=importlib.util.spec_from_file_location('qa4check',ROOT/'scripts/architecture/check_common_qa_v4.py')
    checker=importlib.util.module_from_spec(checker_spec);checker_spec.loader.exec_module(checker)
    with tempfile.TemporaryDirectory(prefix='via-qa-v4-') as temp:
        temp=Path(temp);generated=temp/'reproduced'
        subprocess.run([sys.executable,str(ROOT/'scripts/architecture/build_common_qa_v4.py'),'--output',str(generated)],check=True,stdout=subprocess.DEVNULL)
        for filename in ['ledger.json','summary.json','comparison.csv']:
            assert (generated/filename).read_bytes()==(OUT/'artifacts'/filename).read_bytes(),filename
        for filename in ['02-configurations-and-scenarios.md','03-results-and-ranges.md']:
            assert (generated/'documents'/filename).read_bytes()==(OUT/filename).read_bytes(),filename
        for name in ['specification.json','oracles.json']:shutil.copy(OUT/name,temp/name)
        shutil.copytree(OUT/'artifacts',temp/'artifacts')
        checker.OUT=temp
        mutations=[]
        for tag in ['missing_cost','duplicate_cost','wrong_rate','missing_workspace','wrong_slot','wrong_speech_duration']:
            altered=copy.deepcopy(data)
            if tag=='missing_cost':altered['cases'][0]['usage'].pop()
            elif tag=='duplicate_cost':altered['cases'][0]['usage'].append(copy.deepcopy(altered['cases'][0]['usage'][0]))
            elif tag=='wrong_rate':altered['cases'][0]['usage'][0]['rate_usd']*=2
            elif tag=='missing_workspace':
                c=next(c for c in altered['cases'] if c['config_id']=='DP3-B-R2' and c['case_id']=='U11')
                n=next(n for n in c['nodes'] if n['id']=='now_sam_infer');n['state']={}
                c['memory_intervals']=[i for i in c['memory_intervals'] if i['owner']!='now_sam_infer:sam_workspace']
            elif tag=='wrong_slot':next(n for n in altered['cases'][0]['nodes'] if n['resource']=='semantic')['resource_slot']=99
            else:next(n for n in altered['cases'][0]['nodes'] if n['resource']=='speech')['audio_s']+=1
            (temp/'artifacts/ledger.json').write_text(json.dumps(altered))
            try:
                with contextlib.redirect_stdout(io.StringIO()):checker.main()
            except (AssertionError,KeyError):mutations.append(tag)
            else:raise AssertionError('Undetected mutation: '+tag)
    report=dict(status='PASS',xlsx_sha256=hashlib.sha256(sheetpath.read_bytes()).hexdigest(),
                summary_cells=count,exported_formulas=formula_count,blank_observations=399,
                numeric_and_error_cache_check='PASS',deterministic_ledger_and_documents='PASS',
                negative_validation_probes=mutations,all_sheet_previews=10,
                native_excel='NOT_RUN',model_or_product_run=False)
    (OUT/'artifacts/artifact-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
