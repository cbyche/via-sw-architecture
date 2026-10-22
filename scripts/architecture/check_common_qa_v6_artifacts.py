#!/usr/bin/env python3
"""Read-only validation of new QA07 sheets and unchanged v5 workbook contents."""
import copy
import hashlib
import json
import math
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
P=ROOT/'docs/architecture/12-decisions/decision-packages/current/evaluation'


def main():
    old=P/'common-qa-v5/artifacts/outputs/01a15dfe-v5/VIA-QA-common-v5.xlsx'
    new=P/'common-qa-v6/artifacts/outputs/01a15dfe-v6/VIA-QA-common-v6.xlsx'
    ow=openpyxl.load_workbook(old,data_only=False);nw=openpyxl.load_workbook(new,data_only=False)
    ov=openpyxl.load_workbook(old,data_only=True);nv=openpyxl.load_workbook(new,data_only=True)
    assert nw.sheetnames==ow.sheetnames+['QA07사례','QA07의무']
    allowed={'전체비교':{'A3','A32',*[f'G{r}' for r in range(21,31)]},'의미판정':{'A3'}}
    compared=0
    for name in ow.sheetnames:
        before,after=ow[name],nw[name]
        assert before.freeze_panes==after.freeze_panes,name
        assert set(before.tables)==set(after.tables),name
        for key in before.tables:assert before.tables[key].ref==after.tables[key].ref
        for row in before:
            for cell in row:
                if cell.coordinate in allowed.get(name,set()):continue
                n=after[cell.coordinate];assert cell.value==n.value,(name,cell.coordinate,cell.value,n.value)
                a,b=ov[name][cell.coordinate].value,nv[name][cell.coordinate].value
                if isinstance(a,(int,float)) and isinstance(b,(int,float)):assert math.isclose(a,b,abs_tol=1e-7)
                else:assert a==b,(name,cell.coordinate,a,b)
                if cell.value is not None:
                    assert cell.number_format==n.number_format,(name,cell.coordinate,'format')
                    assert copy.copy(cell.font)==copy.copy(n.font),(name,cell.coordinate,'font')
                compared+=1
    assert nv['QA07사례'].max_row==184
    assert nv['QA07의무'].max_row==1139
    for row in range(6,13):
        assert [nv['QA07사례'].cell(row,c).value for c in range(2,12)]==[24,0,0,24,0,'미산출','미산출','미산출','미산출','미정']
    for row in range(17,185):
        assert nv['QA07사례'].cell(row,12).value=='UNKNOWN'
        registry_row=6+(row-17)%24
        assert nw['QA07사례'].cell(row,14).value==f'=IF(Q{registry_row}="","",Q{registry_row})'
        assert all(nv['QA07사례'].cell(row,c).value in [None,''] for c in [9,10,11,14])
    for row in range(6,1140):
        assert all(nv['QA07의무'].cell(row,c).value in [None,''] for c in range(7,11))
        assert nv['QA07의무'].cell(row,11).value=='UNKNOWN'
    for row in range(21,31):assert nw['전체비교'].cell(row,7).value.startswith("='QA07사례'!")
    for sheet in nv:
        for row in sheet:
            for cell in row:assert cell.data_type!='e',(sheet.title,cell.coordinate,cell.value)
    report=dict(status='PASS',xlsx_sha256=hashlib.sha256(new.read_bytes()).hexdigest(),
                sheets=len(nw.sheetnames),unchanged_parent_cells=compared,draft_case_rows=168,
                obligation_rows=1134,unfilled_observation_cells=4536,
                parent_values_formulas_formats_and_panes='PASS',native_excel='NOT_RUN',actual_measurement=False)
    (P/'common-qa-v6/artifacts/artifact-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report))


if __name__=='__main__':main()
