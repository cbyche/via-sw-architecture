#!/usr/bin/env python3
"""Read-only XLSX validation and deterministic reproduction of the v5 ledger."""
import copy
import contextlib
import hashlib
import importlib.util
import io
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/architecture/12-decisions/decision-packages/current/evaluation/common-qa-v5'
FILES=OUT/'artifacts/outputs/01a15dfe-v5'

def eq(a,b):assert math.isclose(a,b,abs_tol=1e-7),(a,b)

def main():
    d=json.loads((OUT/'artifacts/ledger.json').read_text());book=FILES/'VIA-QA-common-v5.xlsx'
    w=openpyxl.load_workbook(book,data_only=True,read_only=True);f=openpyxl.load_workbook(book,data_only=False,read_only=True)
    meta=json.loads((FILES/'render-verification.json').read_text());assert w.sheetnames==meta['sheets'] and len(w.sheetnames)==11
    assert meta['ledger_sha256']==hashlib.sha256((OUT/'artifacts/ledger.json').read_bytes()).hexdigest()
    configs=['BASE','DP2-B','BASE','DP3-B-R1','DP3-B-R2','DP3-B-R3','BASE','DP5-B','BASE','DP44-B']
    for row,cid in enumerate(configs,6):
        s=d['summary'][cid]
        expected=[s['qa01'],s['qa01_score'],s['wait']['L'],s['wait']['B'],s['wait']['R'],s['qa03'],s['qa03_score']]
        for j,v in enumerate(expected,3):eq(w['전체비교'].cell(row,j).value,v)
        assert w['전체비교'].cell(row+15,3).value==w['전체비교'].cell(row+15,7).value=='미산출'
        for j,q in [(4,'qa04'),(5,'qa05'),(6,'qa06'),(8,'qa08')]:eq(w['전체비교'].cell(row+15,j).value,s[q])
    for ci,cfg in enumerate(d['configs']):
        for gi,group in enumerate(['WAIT_USER','WAIT_PREVIOUS','WAIT_BOTH']):
            r=36+ci*3+gi;s=d['summary'][cfg['id']]['wait_profiles'][group]
            for j,k in [(3,'scheduled'),(4,'onsets'),(5,'withdrawn'),(6,'L'),(7,'B'),(8,'E'),(9,'R')]:
                value=w['전체비교'].cell(r,j).value
                if s[k] is None:assert value=='미산출'
                else:eq(value,s[k])
    for i,c in enumerate(d['cases'],6):
        for j,k in [(5,'goal_wall_s'),(6,'qa03'),(7,'qa05'),(8,'qa06'),(9,'qa08')]:eq(w['목표계산'].cell(i,j).value,c[k])
    all_events=[e for c in d['cases'] for e in c['events']]
    for i,e in enumerate(all_events,6):
        for j,k in [(15,'L'),(16,'B'),(17,'E'),(18,'R'),(19,'D')]:
            v=w['사건계산'].cell(i,j).value
            if e[k] is None:assert v in (None,'')
            else:eq(v,e[k])
    for sheet,col,count in [('지칭판정',5,231),('의미판정',4,168)]:
        assert all(r[0] in (None,'') for r in w[sheet].iter_rows(min_row=6,max_row=5+count,min_col=col,max_col=col,values_only=True))
    for sheet in w:
        for row in sheet:
            for c in row:assert c.data_type!='e',(sheet.title,c.coordinate,c.value)
    formula_count=sum(c.data_type=='f' for sheet in f for row in sheet for c in row)
    w.close();f.close()
    checker_spec=importlib.util.spec_from_file_location('qa5check',ROOT/'scripts/architecture/check_common_qa_v5.py');checker=importlib.util.module_from_spec(checker_spec);checker_spec.loader.exec_module(checker)
    with tempfile.TemporaryDirectory(prefix='via-qa-v5-') as tmp:
        temp=Path(tmp);generated=temp/'reproduced';subprocess.run([sys.executable,str(ROOT/'scripts/architecture/build_common_qa_v5.py'),'--output',str(generated)],check=True,stdout=subprocess.DEVNULL)
        for name in ['ledger.json','summary.json','comparison.csv']:assert (generated/name).read_bytes()==(OUT/'artifacts'/name).read_bytes(),name
        for name in ['02-configurations-and-scenarios.md','03-results-and-ranges.md']:assert (generated/'documents'/name).read_bytes()==(OUT/name).read_bytes(),name
        for name in ['specification.json','oracles.json','preserved-v4-sha256.json']:shutil.copy(OUT/name,temp/name)
        (temp/'artifacts').mkdir();shutil.copy(OUT/'artifacts/comparison.csv',temp/'artifacts/comparison.csv');checker.OUT=temp
        probes=[]
        for tag in ['missing_cost','wrong_received_time','wrong_normal_class','missing_workspace','missing_weight','missing_relation_producer','wrong_slot','wrong_prefix','wrong_score','missing_completion_edge','missing_original_model_budget']:
            a=copy.deepcopy(d)
            if tag=='missing_cost':a['cases'][0]['usage'].pop()
            elif tag=='wrong_received_time':next(e for c in a['cases'] for e in c['events'] if e['kind']=='AGENT_QUESTION')['received_s']+=1
            elif tag=='wrong_normal_class':a['cases'][0]['events'][0]['profile']='WAIT_USER'
            elif tag=='missing_workspace':
                c=next(c for c in a['cases'] if c['config_id']=='DP3-B-R2' and c['case_id']=='U11');c['memory_intervals']=[q for q in c['memory_intervals'] if q['owner']!='now_sam_infer:sam_workspace']
            elif tag=='missing_weight':
                c=next(c for c in a['cases'] if c['config_id']=='DP3-B-R2' and c['case_id']=='U11');c['memory_intervals']=[q for q in c['memory_intervals'] if not q['owner'].startswith('weight:')]
            elif tag=='missing_relation_producer':
                c=next(c for c in a['cases'] if c['config_id']=='DP5-B' and c['case_id']=='U14');next(n for n in c['nodes'] if n['id']=='relation_prior')['state']={};c['memory_intervals']=[q for q in c['memory_intervals'] if q['owner']!='relation_prior:relation_producer']
            elif tag=='wrong_slot':next(n for n in a['cases'][0]['nodes'] if n['resource']=='semantic')['resource_slot']=99
            elif tag=='wrong_prefix':next(n for n in a['cases'][0]['nodes'] if n.get('observes_prefix_of'))['duration_s']+=1
            elif tag=='wrong_score':a['summary']['BASE']['qa01_score']=6
            elif tag=='missing_completion_edge':next(n for n in a['cases'][0]['nodes'] if n['id']=='terminal')['causal_parents']=[]
            else:next(n for c in a['cases'] for n in c['nodes'] if n['id']=='original_work')['tokens_in']=0
            (temp/'artifacts/ledger.json').write_text(json.dumps(a))
            try:
                with contextlib.redirect_stdout(io.StringIO()):checker.main()
            except (AssertionError,KeyError):probes.append(tag)
            else:raise AssertionError('Undetected: '+tag)
    report=dict(status='PASS',xlsx_sha256=hashlib.sha256(book.read_bytes()).hexdigest(),sheets=11,events=294,goals=168,comparison_rows=10,
                exported_formulas=formula_count,blank_accuracy_observations=399,deterministic_reproduction='PASS',negative_validation_probes=probes,
                native_excel='NOT_RUN',product_measurement=False)
    (OUT/'artifacts/artifact-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,indent=2))

if __name__=='__main__':main()
