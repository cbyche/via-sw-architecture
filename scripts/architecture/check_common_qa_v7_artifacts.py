#!/usr/bin/env python3
"""Read-only spreadsheet preservation and v7 result reconciliation."""
import copy
import hashlib
import json
import math
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
P=ROOT/'docs/architecture/12-decisions/decision-packages/current/evaluation'


def main():
    old=P/'common-qa-v6/artifacts/outputs/01a15dfe-v6/VIA-QA-common-v6.xlsx'
    new=P/'common-qa-v7/artifacts/outputs/01a15dfe-v7/VIA-QA-common-v7.xlsx'
    ow=openpyxl.load_workbook(old);nw=openpyxl.load_workbook(new)
    ov=openpyxl.load_workbook(old,data_only=True);nv=openpyxl.load_workbook(new,data_only=True)
    d=json.loads((P/'common-qa-v7/artifacts/ledger.json').read_text());ap=d['appendices']
    assert nw.sheetnames==ow.sheetnames+['QA04변경','QA상세부록']
    allowed={'전체비교':{'A3','D20','E20','H20',*[f'{c}{r}' for c in ['D','H'] for r in range(21,31)]},'변경과제':{'A2','A3'}}
    compared=0
    for name in ow.sheetnames:
        before,after=ow[name],nw[name]
        assert before.freeze_panes==after.freeze_panes,name
        assert set(before.tables)==set(after.tables),name
        assert list(before.data_validations)==list(after.data_validations),name
        for key in before.tables:assert before.tables[key].ref==after.tables[key].ref
        for row in before:
            for cell in row:
                if cell.coordinate in allowed.get(name,set()):continue
                n=after[cell.coordinate];assert cell.value==n.value,(name,cell.coordinate,cell.value,n.value)
                a,b=ov[name][cell.coordinate].value,nv[name][cell.coordinate].value
                if isinstance(a,(int,float)) and isinstance(b,(int,float)):assert math.isclose(a,b,abs_tol=1e-7),(name,cell.coordinate,a,b)
                else:assert a==b,(name,cell.coordinate,a,b)
                if cell.value is not None:
                    assert cell.number_format==n.number_format,(name,cell.coordinate,'format')
                    for attr in ['font','fill','alignment']:assert copy.copy(getattr(cell,attr))==copy.copy(getattr(n,attr)),(name,cell.coordinate,attr)
                compared+=1
    for r in range(6,13):assert [nv['QA04변경'].cell(r,c).value for c in range(2,7)]==[8,0,'미측정','미산출','실제 patch 없음']
    for r in range(17,73):
        assert nv['QA04변경'].cell(r,13).value=='NOT_MEASURED'
        assert all(nv['QA04변경'].cell(r,c).value in [None,''] for c in range(7,13))
    cmp=['BASE','DP2-B','BASE','DP3-B-R1','DP3-B-R2','DP3-B-R3','BASE','DP5-B','BASE','DP44-B']
    for i,cfg in enumerate(cmp):
        r=21+i;assert nv['전체비교'].cell(r,4).value=='미측정'
        assert nw['전체비교'].cell(r,4).value.startswith("='QA04변경'!")
        assert math.isclose(nv['전체비교'].cell(r,8).value,d['summary'][cfg]['qa08_gb'],abs_tol=1e-9)
    s=nv['QA상세부록'];row=6
    groups=[('QA01',ap['qa01_types'],'mean_s'),('QA02',ap['qa02_tags'],None),('QA05',ap['qa05_buckets'],'usd'),('QA07',ap['qa07_types'],None)]
    for qa,items,key in groups:
        for item in items:
            assert s.cell(row,1).value==qa and s.cell(row,2).value==item['config_id']
            v=s.cell(row,7).value;expected=item.get(key) if key else None
            if expected is None:assert v=='미산출',(row,v)
            else:assert math.isclose(v,expected,abs_tol=1e-8),(row,v,expected)
            assert s.cell(row,9).value==0;row+=1
    for a in ap['memory_peaks']:
        expected=a['value_mb']/(1000 if a['qa']=='qa08' else 1)
        assert math.isclose(s.cell(row,7).value,expected,abs_tol=1e-8);row+=1
    for sheet in nv:
        for cells in sheet:
            for cell in cells:assert cell.data_type!='e',(sheet.title,cell.coordinate,cell.value)
    result=dict(status='PASS',xlsx_sha256=hashlib.sha256(new.read_bytes()).hexdigest(),sheets=len(nw.sheetnames),
                unchanged_parent_cells=compared,parent_values_formulas_styles_validations='PASS',
                qa04_unmeasured=56,appendix_reconciliation='PASS',native_excel='NOT_RUN',actual_measurement=False)
    (P/'common-qa-v7/artifacts/artifact-validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');print(json.dumps(result))


if __name__=='__main__':main()
